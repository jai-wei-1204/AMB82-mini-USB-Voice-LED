from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "ProjectReport.md"
OUTPUT = ROOT.parent / "output" / "docx" / "AMB82_mini_語音方向燈專題報告.docx"

FONT = "Microsoft JhengHei"
ACCENT = "1F4E79"
LIGHT_BLUE = "D9EAF7"
BORDER = "D9D9D9"


def set_font(run, size=None, bold=None, color=None, name=FONT):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def border_cell(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        tag = OxmlElement(f"w:{side}")
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "4")
        tag.set(qn("w:color"), BORDER)
        borders.append(tag)
    tc_pr.append(borders)


def set_cell_text(cell, text, header=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if header else WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(text.strip())
    set_font(r, 9 if header else 8.5, bold=header, color="FFFFFF" if header else "000000")
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if header:
        shade(cell, ACCENT)
    border_cell(cell)


def add_table(doc, rows):
    if len(rows) < 2:
        return
    data = [[c.strip() for c in row.strip().strip("|").split("|")] for row in rows]
    data = [r for r in data if not all(set(c.replace("-", "").replace(":", "")) == set() for c in r)]
    cols = max(len(r) for r in data)
    table = doc.add_table(rows=0, cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for row_no, values in enumerate(data):
        cells = table.add_row().cells
        for col_no in range(cols):
            set_cell_text(cells[col_no], values[col_no] if col_no < len(values) else "", header=row_no == 0)
    doc.add_paragraph().paragraph_format.space_after = Pt(3)


def add_body(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.35
    r = p.add_run(text)
    set_font(r, 10.5)


def add_code(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(text)
    set_font(r, 8, name="Consolas")
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F2F2F2")
    p_pr.append(shd)


def add_heading(doc, level, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14 if level == 2 else 10)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    if level == 1:
        p.style = "Title"
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_font(r, 22, bold=True)
    elif level == 2:
        set_font(r, 15, bold=True)
    else:
        set_font(r, 12, bold=True)


def add_footer(section):
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run("AMB82-mini USB 語音方向燈專題報告")
    set_font(r, 8, color="666666")
    footer.add_run("  |  ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)


def main():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    add_footer(section)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(10.5)
    doc.core_properties.title = "AMB82-mini USB 語音方向燈專題報告"

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    i = 0
    in_code = False
    code_lines = []
    while i < len(lines):
        line = lines[i]
        if line.startswith("```"):
            if in_code:
                add_code(doc, "\n".join(code_lines))
                code_lines = []
            in_code = not in_code
            i += 1
            continue
        if in_code:
            code_lines.append(line)
            i += 1
            continue
        if line.startswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].startswith("|"):
                table_lines.append(lines[i])
                i += 1
            add_table(doc, table_lines)
            continue
        if line.startswith("# "):
            add_heading(doc, 1, line[2:])
        elif line.startswith("## "):
            add_heading(doc, 2, line[3:])
        elif line.startswith("### "):
            add_heading(doc, 3, line[4:])
        elif line.startswith("> "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(line[2:])
            set_font(r, 10, color="666666")
        elif line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(3)
            r = p.add_run(line[2:])
            set_font(r, 10.5)
        elif line and line[0].isdigit() and ". " in line[:4]:
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.space_after = Pt(3)
            r = p.add_run(line.split(". ", 1)[1])
            set_font(r, 10.5)
        elif line.strip():
            add_body(doc, line)
        i += 1

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
